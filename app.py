"""Beginner-friendly PDF to Word converter with optional Gemini heading detection."""

import hashlib
import io
import json
import os
from pathlib import Path

import pymupdf
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from dotenv import load_dotenv
from google import genai
from google.genai import types
from pydantic import BaseModel

load_dotenv()
MODEL = "gemini-3.5-flash-lite"


class HeadingResponse(BaseModel):
    heading_ids: list[int]


def extract_pdf_blocks(pdf_bytes: bytes):
    """Extract readable text blocks, preserving original words and page order."""
    blocks = []
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as pdf:
        if pdf.needs_pass:
            raise ValueError("This PDF is password-protected. Please unlock it first.")

        number_of_pages = len(pdf)
        for page_number, page in enumerate(pdf, start=1):
            for block in page.get_text("blocks", sort=True):
                # Block type 0 means text; ignore image metadata, if present.
                if len(block) > 6 and block[6] != 0:
                    continue

                # Join wrapped text lines into a normal Word paragraph.
                lines = [line.strip() for line in block[4].splitlines() if line.strip()]
                text = " ".join(lines)
                if text:
                    blocks.append({
                        "id": len(blocks),
                        "page": page_number,
                        "text": text,
                    })

    return blocks, number_of_pages


def identify_headings(blocks, api_key: str):
    """Ask Gemini only for block IDs. The AI never rewrites PDF text."""
    client = genai.Client(api_key=api_key)
    heading_ids = set()

    # Small batches keep prompts inexpensive and make the answer predictable.
    for start in range(0, len(blocks), 35):
        batch = blocks[start:start + 35]
        samples = [{"id": b["id"], "text": b["text"][:250]} for b in batch]
        prompt = (
            "You are classifying text blocks extracted from a PDF. "
            "Identify ONLY clear document titles or section/subsection headings. "
            "Do NOT select ordinary paragraphs, bullet points, or sentences. "
            "Be conservative. Return heading IDs only. Never rewrite text.\n\n"
            + json.dumps(samples, ensure_ascii=False)
        )
        response = client.models.generate_content(
            model=MODEL,
            contents=prompt,
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=HeadingResponse,
                temperature=0,
            ),
        )
        if not response.text:
            raise RuntimeError("Gemini did not return a usable response.")

        result = json.loads(response.text)
        valid_ids = {b["id"] for b in batch}
        for block_id in result.get("heading_ids", []):
            if isinstance(block_id, int) and block_id in valid_ids:
                heading_ids.add(block_id)

    return heading_ids


def make_word_document(blocks, heading_ids):
    """Generate the editable Word file, using the locally extracted text."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.space_after = Pt(6)

    for block in blocks:
        if block["id"] in heading_ids:
            document.add_heading(block["text"], level=1)
        else:
            document.add_paragraph(block["text"])

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def main():
    st.set_page_config(page_title="AI PDF to Word Converter", page_icon="📄")
    st.title("📄 AI PDF to Word Converter")
    st.write("Upload a PDF containing selectable text and download an editable Word file.")

    pdf_file = st.file_uploader("Upload a text-based PDF", type=["pdf"])
    use_ai = st.checkbox("Use Gemini AI to identify headings", value=True)
    if not use_ai:
        st.info("Basic mode: the converter works locally without an AI API key.")

    if pdf_file is not None:
        file_bytes = pdf_file.getvalue()
        result_key = hashlib.sha256(file_bytes + str(use_ai).encode()).hexdigest()
        st.caption(f"Selected file: {pdf_file.name}")

        if st.button("Convert to Word", type="primary"):
            try:
                with st.spinner("Extracting text and creating the Word document..."):
                    blocks, page_count = extract_pdf_blocks(file_bytes)
                    if not blocks:
                        raise ValueError(
                            "No selectable text found. This demo does not support scanned/image PDFs."
                        )

                    heading_ids = set()
                    if use_ai:
                        api_key = os.getenv("GEMINI_API_KEY", "").strip()
                        if not api_key:
                            raise ValueError(
                                "GEMINI_API_KEY is missing. Add it to .env, or uncheck the AI option."
                            )
                        heading_ids = identify_headings(blocks, api_key)

                    st.session_state["docx_bytes"] = make_word_document(
                        blocks, heading_ids
                    )
                    st.session_state["result_key"] = result_key
                    st.session_state["output_name"] = Path(pdf_file.name).stem + ".docx"
                    st.session_state["page_count"] = page_count
                    st.session_state["ai_used"] = use_ai
            except Exception as exc:
                st.error(f"Conversion failed: {exc}")

        if st.session_state.get("result_key") == result_key:
            st.success(
                f"Converted {st.session_state['page_count']} page(s) successfully"
                + (" with Gemini heading detection." if st.session_state["ai_used"] else ".")
            )
            st.download_button(
                label="Download Word document",
                data=st.session_state["docx_bytes"],
                file_name=st.session_state["output_name"],
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

    st.caption("Text-only PDFs. No OCR. Complex formatting may not match the PDF exactly.")


if __name__ == "__main__":
    main()
